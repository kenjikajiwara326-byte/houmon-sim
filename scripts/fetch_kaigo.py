#!/usr/bin/env python3
"""Fetch & clean 介護保険 statistics for the 訪問診療 demand model.

Sources (all e-Stat file downloads, no API key):

A. 介護保険事業状況報告 年報 令和6年度 (FY2024; 政府統計コード 00450351, 公開 2026-08-27)
   都道府県別 (…T) and 保険者別 (…h) tables:
     04-1-1  要介護（要支援）認定者数 男女計 総数 (年度末 = 2025-03-31 現在, 第1号+第2号)
     05-2    居宅介護（介護予防）サービスのサービス別受給者数【現物給付】(当年度累計 = 延人月)
     06-2    地域密着型（介護予防）サービスのサービス別受給者数【現物給付】(延人月)
     07-1    施設介護サービス受給者数 (延人月)
   累計 tables cover service months 2024-03 .. 2025-02; we divide by 12 -> average monthly persons.

B. 介護給付費等実態統計 令和6年度 年報 (00450049; 審査月 2024-05 .. 2025-04 = service months 2024-04 .. 2025-03)
     表35/36 件数，都道府県・(介護予防)サービス種類・要介護(要支援)状態区分別
     表41    単位数・回数・日数・件数，介護予防サービス種類内容・要支援状態区分別  (national)
     表45    回数・日数，介護サービス(居宅サービス等)種類内容・要介護状態区分別 (national)
     表60    単位数・回数・件数，サービスコード別 (national)
   Used to get the physician (医師) share of 居宅療養管理指導 by care level (national only; no
   prefecture x 職種 table exists).

C. 住民基本台帳に基づく人口… 2025年 表25-03 市区町村別 (00200241) -> 団体コード for name matching.

Outputs (data/work/):
  kaigo_hokensha.csv, kyotaku_ryoyo_pref.csv, nintei_pref.csv, util_rates.csv, phys_ratio_national.csv
"""
import io
import os
import re
import subprocess
import sys
import time
import unicodedata

import numpy as np
import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "raw", "kaigo")
WORK = os.path.join(ROOT, "data", "work")
DL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId={}&fileKind={}"

# (local filename, statInfId, fileKind)
FILES = {
    # 介護保険事業状況報告 年報 令和6年度
    "jigyo_R6_04-1-1T.xlsx": ("000040491597", 0),
    "jigyo_R6_05-2T.xlsx": ("000040491607", 0),
    "jigyo_R6_06-2T.xlsx": ("000040491610", 0),
    "jigyo_R6_07-1T.xlsx": ("000040491612", 0),
    "jigyo_R6_04-1-1h.xlsx": ("000040491652", 0),
    "jigyo_R6_05-2h.xlsx": ("000040491662", 0),
    "jigyo_R6_06-2h.xlsx": ("000040491665", 0),
    "jigyo_R6_07-1h.xlsx": ("000040491667", 0),
    # 介護給付費等実態統計 令和6年度
    "jittai_fy2024_t35.csv": ("000040332328", 1),
    "jittai_fy2024_t36.csv": ("000040332329", 1),
    "jittai_fy2024_t41.csv": ("000040332334", 1),
    "jittai_fy2024_t45.csv": ("000040332338", 1),
    "jittai_fy2024_t60.xlsx": ("000040332355", 0),
    # 住民基本台帳 2025 市区町村別 (団体コード)
    "jukiban_2025_25-03.xlsx": ("000040306653", 0),
}

LEVELS = ["shien1", "shien2", "kaigo1", "kaigo2", "kaigo3", "kaigo4", "kaigo5"]
KAIGO = LEVELS[2:]
PREFS = ["北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県", "茨城県", "栃木県", "群馬県",
         "埼玉県", "千葉県", "東京都", "神奈川県", "新潟県", "富山県", "石川県", "福井県", "山梨県", "長野県",
         "岐阜県", "静岡県", "愛知県", "三重県", "滋賀県", "京都府", "大阪府", "兵庫県", "奈良県", "和歌山県",
         "鳥取県", "島根県", "岡山県", "広島県", "山口県", "徳島県", "香川県", "愛媛県", "高知県", "福岡県",
         "佐賀県", "長崎県", "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県"]
PREF_CODE = {p: f"{i + 1:02d}" for i, p in enumerate(PREFS)}
KANTO = ["茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県"]
# 一部事務組合・広域連合 保険者 in Kanto -> member municipality codes (FY2024)
KOUIKI_MEMBERS = {("埼玉県", "大里広域市町村圏組合"): ["11202", "11218", "11408"]}  # 熊谷市・深谷市・寄居町


def download():
    os.makedirs(RAW, exist_ok=True)
    for fn, (sid, kind) in FILES.items():
        path = os.path.join(RAW, fn)
        if os.path.exists(path) and os.path.getsize(path) > 0:
            continue
        url = DL.format(sid, kind)
        print("download", fn, url, file=sys.stderr)
        subprocess.run(["curl", "-sS", "-L", "--fail", "-o", path, url], check=True)
        time.sleep(0.5)


def num(x):
    if x is None:
        return np.nan
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(",", "")
    if s in ("-", "", "…", "x", "X", "***"):
        return 0.0 if s == "-" else np.nan
    try:
        return float(s)
    except ValueError:
        return np.nan


def norm(s):
    s = unicodedata.normalize("NFKC", str(s)).replace(" ", "").replace("　", "")
    return s.replace("ヶ", "ケ").replace("ヵ", "カ")


# ---------------------------------------------------------------- 年報 parsers
def read_sheet(fn, sheet_idx=0):
    wb = openpyxl.load_workbook(os.path.join(RAW, fn), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[sheet_idx]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def parse_h_simple(fn, sheet_idx):
    """保険者別 table with one 7-level block per sheet (04-1-1h, 07-1h)."""
    rows = read_sheet(fn, sheet_idx)
    out = []
    for r in rows:
        if r[0] in PREFS and r[1]:
            out.append([r[0], r[1]] + [num(v) for v in r[2:9]])
    return pd.DataFrame(out, columns=["pref", "name"] + LEVELS)


def parse_h_services(fn, services):
    """保険者別 service x level table (05-2h, 06-2h): header row has service names, next row levels."""
    rows = read_sheet(fn, 0)
    hdr = next(i for i, r in enumerate(rows) if r[1] in ("保険者名", "保険者"))
    names = rows[hdr]
    res = {}
    for svc in services:
        c = next(j for j, v in enumerate(names) if v and norm(v) == norm(svc))
        recs = []
        for r in rows[hdr + 2:]:
            if r[0] in PREFS and r[1]:
                recs.append([r[0], r[1]] + [num(v) for v in r[c:c + 8]])
        res[svc] = pd.DataFrame(recs, columns=["pref", "name"] + LEVELS + ["total"])
    return res


def parse_T_blocks(fn, sheet_idx=0):
    """都道府県別 table laid out as repeated 9-col blocks [都道府県, 7 levels, 合計].
    Returns {block title: DataFrame(pref, levels..., total)}."""
    rows = read_sheet(fn, sheet_idx)
    lvl_row = next(i for i, r in enumerate(rows) if "要支援１" in [str(v) for v in r])
    title_row = rows[lvl_row - 1]
    out = {}
    for j, v in enumerate(rows[lvl_row]):
        if v == "要支援１":
            start = j
            title = title_row[start] or rows[lvl_row - 1][start - 1] or f"block{start}"
            recs = []
            for r in rows[lvl_row + 1:]:
                p = r[start - 1]
                if p in PREFS or (p and str(p).replace("　", "").replace(" ", "") in ("全国", "全国計")):
                    p = p if p in PREFS else "全国"
                    recs.append([p] + [num(x) for x in r[start:start + 8]])
            out[norm(title)] = pd.DataFrame(recs, columns=["pref"] + LEVELS + ["total"])
    return out


# ---------------------------------------------------------------- 実態統計 parsers
def read_csv_sjis(fn):
    raw = open(os.path.join(RAW, fn), "rb").read().decode("cp932", errors="replace")
    return [line.split(",") for line in raw.splitlines()]


def jittai_pref_kensu(fn, svc, levels):
    """表35/36: pref x level 件数 (千件) for one service."""
    recs = []
    for r in read_csv_sjis(fn):
        if len(r) > 5 and r[2] == svc and r[3] == "" and r[0] in PREFS + ["全国"]:
            vals = [num(x) * 1000 for x in r[6:6 + len(levels)]]
            recs.append([r[0]] + vals)
    return pd.DataFrame(recs, columns=["pref"] + levels).drop_duplicates("pref")


def jittai_content_counts(fn, svc, col0, levels):
    """表41/45: national 回数 (千回) by content item x level; col0 = first level column index."""
    out = {}
    for r in read_csv_sjis(fn):
        if len(r) > col0 and r[0] == svc and r[1]:
            out[r[1]] = [num(x) * 1000 for x in r[col0:col0 + len(levels)]]
    return pd.DataFrame(out, index=levels).T


def phys_ratio_national():
    """National physician 居宅療養管理指導 'persons' (≈ 件数 of physician service codes) by level.

    表60 gives 回数 and 件数 per service code (annual, all levels) -> visits per claim v_code.
    表45 / 表41 give 回数 per content item x care level. physician 件_L = Σ_code 回数_{code,L} / v_code.
    """
    wb = openpyxl.load_workbook(os.path.join(RAW, "jittai_fy2024_t60.xlsx"), read_only=True)
    codes = {}
    for r in wb[wb.sheetnames[0]].iter_rows(values_only=True):
        if isinstance(r[0], (int, float)) or (r[0] and str(r[0]).isdigit()):
            codes[int(r[0])] = (r[1], num(r[3]), num(r[4]))  # name, 回数(千), 件数(千)
    # content item label -> service code
    item2code_k = {"医師（Ⅰ）（一）": 311111, "医師（Ⅱ）（一）": 311112, "医師（Ⅰ）（二）": 311113,
                   "医師（Ⅱ）（二）": 311114, "医師（Ⅰ）（三）": 311115, "医師（Ⅱ）（三）": 311116}
    item2code_y = {k: v + 30000 for k, v in item2code_k.items()}  # 3411xx 予防
    bldg = {"（一）": "b1", "（二）": "b2_9", "（三）": "b10plus"}

    k = jittai_content_counts("jittai_fy2024_t45.csv", "居宅療養管理指導", 4, KAIGO)
    y = jittai_content_counts("jittai_fy2024_t41.csv", "介護予防居宅療養管理指導", 7, LEVELS[:2])
    rows = []
    for tab, i2c, lv in ((y, item2code_y, LEVELS[:2]), (k, item2code_k, KAIGO)):
        for item, code in i2c.items():
            name, kaisu, kensu = codes[code]
            v = kaisu / kensu
            for L in lv:
                rows.append(dict(level=L, item=item, code=code, bldg=[b for s, b in bldg.items() if s in item][0],
                                 kaisu=tab.loc[item, L], visits_per_claim=v, phys_claims=tab.loc[item, L] / v))
    det = pd.DataFrame(rows)
    agg = det.groupby("level").agg(phys_claims_year=("phys_claims", "sum"), phys_visits_year=("kaisu", "sum"))
    for b in ["b1", "b2_9", "b10plus"]:
        agg[f"share_{b}"] = det[det.bldg == b].groupby("level").phys_claims.sum() / agg.phys_claims_year
    return agg.reindex(LEVELS), det


# ---------------------------------------------------------------- build outputs
def muni_codes():
    wb = openpyxl.load_workbook(os.path.join(RAW, "jukiban_2025_25-03.xlsx"), read_only=True)
    m = {}
    for r in wb[wb.sheetnames[0]].iter_rows(min_row=7, values_only=True):
        if r[0] and r[2] and r[2] != "-" and r[1] in PREFS:
            nm = norm(r[2])
            nm = re.sub(r"^.+?郡(?=.+[町村]$)", "", nm)  # strip 郡 prefix
            m[(r[1], nm)] = str(r[0])[:5]
    return m


def build_hokensha():
    codes = muni_codes()
    nintei = parse_h_simple("jigyo_R6_04-1-1h.xlsx", 0)  # ① 総数 (1号+2号)
    shis = parse_h_simple("jigyo_R6_07-1h.xlsx", 0)      # ① 総数 of 4 facility types (deduplicated)
    parts = {n: parse_h_simple("jigyo_R6_07-1h.xlsx", i) for n, i in
             (("rofuku", 1), ("roken", 4), ("ryoyo", 7), ("iryoin", 10))}
    s5 = parse_h_services("jigyo_R6_05-2h.xlsx", ["居宅療養管理指導", "特定施設入居者生活介護"])
    s6 = parse_h_services("jigyo_R6_06-2h.xlsx", ["認知症対応型共同生活介護", "地域密着型特定施設入居者生活介護",
                                                   "地域密着型介護老人福祉施設入所者生活介護"])
    df = nintei[nintei.pref.isin(KANTO)].rename(columns={L: f"nintei_{L}" for L in LEVELS})
    df["nintei_total"] = df[[f"nintei_{L}" for L in LEVELS]].sum(axis=1)
    key = ["pref", "name"]

    def add(d, prefix, levels, scale=1 / 12):
        d = d[key + levels].copy()
        d[levels] = d[levels] * scale
        return df.merge(d.rename(columns={L: f"{prefix}_{L}" for L in levels}), on=key, how="left")

    df = add(shis, "shisetsu3", KAIGO)
    df = add(s6["地域密着型介護老人福祉施設入所者生活介護"], "chiiki_rofuku", KAIGO)
    for L in KAIGO:
        df[f"shisetsu_{L}"] = df[f"shisetsu3_{L}"] + df[f"chiiki_rofuku_{L}"]
    for n, d in parts.items():
        df = df.merge(d[key].assign(**{f"{n}_total": d[KAIGO].sum(axis=1) / 12}), on=key, how="left")
    tok = s5["特定施設入居者生活介護"][key + ["total"]].rename(columns={"total": "t1"}).merge(
        s6["地域密着型特定施設入居者生活介護"][key + ["total"]].rename(columns={"total": "t2"}), on=key)
    tok["tokutei"] = (tok.t1 + tok.t2) / 12
    df = df.merge(tok[key + ["tokutei"]], on=key, how="left")
    gh = s6["認知症対応型共同生活介護"]
    df = df.merge(gh[key].assign(gh=gh.total / 12), on=key, how="left")
    df = add(s5["居宅療養管理指導"], "kyoryo_all", LEVELS)

    df.insert(0, "pref_code", df.pref.map(PREF_CODE))
    df.insert(1, "code", [codes.get((p, norm(n))) for p, n in zip(df.pref, df.name)])
    df.insert(4, "is_kouiki", [(p, n) in KOUIKI_MEMBERS for p, n in zip(df.pref, df.name)])
    df.insert(5, "member_codes", ["|".join(KOUIKI_MEMBERS.get((p, n), [])) for p, n in zip(df.pref, df.name)])
    unmatched = df[df.code.isna() & ~df.is_kouiki]
    if len(unmatched):
        print("WARNING unmatched 保険者:", unmatched[key].values.tolist(), file=sys.stderr)
    return df.round(2)


def build_pref():
    nin = parse_T_blocks("jigyo_R6_04-1-1T.xlsx", 0)
    nin = list(nin.values())[0]
    shis = list(parse_T_blocks("jigyo_R6_07-1T.xlsx", 0).values())[0]
    s5 = parse_T_blocks("jigyo_R6_05-2T.xlsx", 0)
    s6 = parse_T_blocks("jigyo_R6_06-2T.xlsx", 0)
    kr = s5["居宅療養管理指導"]
    crof = s6[norm("地域密着型介護老人福祉施設入所者生活介護")]
    tok1, tok2 = s5["特定施設入居者生活介護"], s6[norm("地域密着型特定施設入居者生活介護")]
    gh = s6[norm("認知症対応型共同生活介護")]

    def long(d, col, scale=1.0, levels=LEVELS):
        d = d[d.pref.isin(PREFS)].melt(id_vars="pref", value_vars=levels, var_name="level", value_name=col)
        d[col] *= scale
        return d

    nt = long(nin, "nintei")
    nt = nt.merge(long(shis, "shisetsu3", 1 / 12), on=["pref", "level"])
    nt = nt.merge(long(crof, "chiiki_rofuku", 1 / 12), on=["pref", "level"])
    nt["shisetsu"] = nt.shisetsu3 + nt.chiiki_rofuku
    nt = nt.merge(long(tok1, "tokutei_a", 1 / 12), on=["pref", "level"]).merge(
        long(tok2, "tokutei_b", 1 / 12), on=["pref", "level"])
    nt["tokutei"] = nt.tokutei_a + nt.tokutei_b
    nt = nt.drop(columns=["tokutei_a", "tokutei_b"]).merge(long(gh, "gh", 1 / 12), on=["pref", "level"])
    nt["zaitaku"] = nt.nintei - nt.shisetsu
    nt.insert(0, "pref_code", nt.pref.map(PREF_CODE))

    # --- 居宅療養管理指導
    ratio, det = phys_ratio_national()
    kr_l = long(kr, "recipients_all", 1 / 12)  # 年報 05-2T, avg monthly persons (all 職種)
    nat_recip = kr_l.groupby("level").recipients_all.sum()
    ratio["phys_claims_month"] = ratio.phys_claims_year / 12
    ratio["nat_recipients_all_month"] = nat_recip
    ratio["r_phys_per_recipient"] = ratio.phys_claims_month / ratio.nat_recipients_all_month
    kk = jittai_pref_kensu("jittai_fy2024_t36.csv", "居宅療養管理指導", KAIGO)
    ky = jittai_pref_kensu("jittai_fy2024_t35.csv", "介護予防居宅療養管理指導", LEVELS[:2])
    ken = ky.merge(kk, on="pref")
    nat_ken = ken[ken.pref == "全国"].iloc[0]
    ratio["nat_claims_all_month"] = [nat_ken[L] / 12 for L in LEVELS]
    ratio["phys_share_of_claims"] = ratio.phys_claims_month / ratio.nat_claims_all_month
    ken_l = long(ken, "claims_all", 1 / 12)

    kp = kr_l.merge(ken_l, on=["pref", "level"])
    kp = kp.merge(ratio[["r_phys_per_recipient", "phys_share_of_claims", "share_b1", "share_b2_9",
                         "share_b10plus"]], left_on="level", right_index=True)
    kp["recipients"] = kp.recipients_all * kp.r_phys_per_recipient          # primary physician estimate
    kp["recipients_alt_claims"] = kp.claims_all * kp.phys_share_of_claims   # alt: 件数-based
    for b in ["b1", "b2_9", "b10plus"]:
        kp[f"recipients_{b}"] = kp.recipients * kp[f"share_{b}"]
    kp.insert(0, "pref_code", kp.pref.map(PREF_CODE))
    kp = kp[["pref_code", "pref", "level", "recipients", "recipients_all", "claims_all", "recipients_alt_claims",
             "r_phys_per_recipient", "phys_share_of_claims", "recipients_b1", "recipients_b2_9",
             "recipients_b10plus"]]
    for d in (nt, kp):
        d["lv"] = d.level.map(LEVELS.index)
    nt = nt.sort_values(["pref_code", "lv"]).drop(columns="lv")
    kp = kp.sort_values(["pref_code", "lv"]).drop(columns="lv")
    return nt, kp, ratio, det


def util_rates(nt, kp):
    d = kp[["pref_code", "pref", "level", "recipients", "recipients_all"]].merge(
        nt[["pref_code", "level", "nintei", "shisetsu", "zaitaku"]], on=["pref_code", "level"])
    d["u"] = d.recipients / d.zaitaku
    d["u_all_shokushu"] = d.recipients_all / d.zaitaku
    nat = d.groupby("level")[["recipients", "recipients_all", "nintei", "shisetsu", "zaitaku"]].sum().reset_index()
    nat["u"] = nat.recipients / nat.zaitaku
    nat["u_all_shokushu"] = nat.recipients_all / nat.zaitaku
    nat["pref_code"], nat["pref"] = "00", "全国"
    p75 = d.groupby("level")[["u", "u_all_shokushu"]].quantile(0.75).reset_index()
    p75["pref_code"], p75["pref"] = "P75", "p75(47都道府県)"
    out = pd.concat([d, nat, p75], ignore_index=True)
    out["level"] = pd.Categorical(out.level, LEVELS, ordered=True)
    return out.sort_values(["pref_code", "level"])[["pref_code", "pref", "level", "recipients", "recipients_all",
                                                    "nintei", "shisetsu", "zaitaku", "u", "u_all_shokushu"]]


def main():
    download()
    os.makedirs(WORK, exist_ok=True)
    h = build_hokensha()
    h.to_csv(os.path.join(WORK, "kaigo_hokensha.csv"), index=False)
    nt, kp, ratio, det = build_pref()
    nt.round(2).to_csv(os.path.join(WORK, "nintei_pref.csv"), index=False)
    kp.round(4).to_csv(os.path.join(WORK, "kyotaku_ryoyo_pref.csv"), index=False)
    ratio.round(4).to_csv(os.path.join(WORK, "phys_ratio_national.csv"))
    u = util_rates(nt, kp)
    u.round(5).to_csv(os.path.join(WORK, "util_rates.csv"), index=False)

    print(f"kaigo_hokensha.csv rows={len(h)}  kyotaku_ryoyo_pref.csv rows={len(kp)}  "
          f"nintei_pref.csv rows={len(nt)}  util_rates.csv rows={len(u)}")
    print("\nNational physician ratio by level:\n", ratio.round(3).to_string())
    show = u[u.pref.isin(["全国", "p75(47都道府県)"] + KANTO)]
    piv = show.pivot_table(index=["pref_code", "pref"], columns="level", values="u", observed=True)
    pd.set_option("display.width", 200)
    print("\nu = physician 居宅療養管理指導 recipients / (認定者 - 施設受給者):\n", (piv * 100).round(1).to_string(),
          "  (%)")


if __name__ == "__main__":
    main()
